# one-off dump of source xlsx — UTF-8 files
from openpyxl import load_workbook
from pathlib import Path

out_dir = Path(r"C:\Users\MiniStation\SynologyDrive\dev\Prod_Plan_MONZA")

files = [
    Path(r"C:\Users\MiniStation\SynologyDrive\dev\Prod_Plan_MONZA")
    / "Монитор распределения заказов на Мебельное производство 07.09.2026.xlsx",
    Path(r"C:\Users\MiniStation\SynologyDrive\dev\Fasad_Monzana")
    / "ТЗ Фасады калькулятор 07.09.2026.xlsx",
]

for path in files:
    lines = []
    lines.append("=" * 80)
    lines.append(f"FILE: {path.name}")
    lines.append("=" * 80)
    for data_only in (True, False):
        wb = load_workbook(path, data_only=data_only)
        lines.append(f"\n### data_only={data_only}  sheets={wb.sheetnames}")
        for name in wb.sheetnames:
            ws = wb[name]
            lines.append("-" * 80)
            lines.append(f"SHEET: {name!r}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
            merges = [str(m) for m in ws.merged_cells.ranges]
            if merges:
                lines.append("MERGES: " + ", ".join(merges))
            comments = []
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=ws.max_column or 1):
                cells = []
                for c in row:
                    if c.comment and c.comment.text:
                        comments.append(f"{c.coordinate}: {c.comment.text}")
                    v = c.value
                    if v is None:
                        continue
                    cells.append(f"{c.coordinate}={v!r}")
                if cells:
                    lines.append(" | ".join(cells))
            if comments:
                lines.append("COMMENTS:")
                lines.extend(comments)
    out = out_dir / (path.stem[:40] + ".dump.txt")
    out.write_text("\n".join(lines), encoding="utf-8")
    print("wrote", out, "chars", sum(len(x) for x in lines))
